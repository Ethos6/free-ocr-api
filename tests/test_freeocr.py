import io
import json
import subprocess
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "examples"))

from make_samples import make_all  # noqa: E402

from freeocr import extract_document  # noqa: E402
from freeocr.export import export  # noqa: E402
from freeocr.ocr import available_engines  # noqa: E402
from freeocr.parsers import find_dates, find_entities, key_values, parse_amount, redact  # noqa: E402

needs_ocr = pytest.mark.skipif(not available_engines(), reason="no OCR engine installed")


@pytest.fixture(scope="session")
def samples(tmp_path_factory):
    return make_all(tmp_path_factory.mktemp("samples"))


# ---------------------------------------------------------------- parsers

@pytest.mark.parametrize("raw,val", [
    ("$1,234.56", 1234.56), ("1.234,56 €", 1234.56), ("₹1,23,456.00", 123456.0),
    ("(12.00)", -12.0), ("12,50", 12.5), ("GBP 345.60", 345.6), ("abc", None),
])
def test_parse_amount(raw, val):
    assert parse_amount(raw) == val


def test_dates():
    iso = [d[2] for d in find_dates("15/03/2026, March 30, 2026, 2026-04-01, 5th Jan 2027")]
    assert iso == ["2026-03-15", "2026-03-30", "2026-04-01", "2027-01-05"]
    assert find_dates("03/04/2026", dayfirst=True)[0][2] == "2026-04-03"
    assert find_dates("03/04/2026", dayfirst=False)[0][2] == "2026-03-04"


def test_entities_and_redact():
    t = "mail a@b.com call +1 (555) 123-4567 IBAN DE89 3704 0044 0532 0130 00 card 4111 1111 1111 1111 GSTIN 29ABCDE1234F1Z5"
    e = find_entities(t)
    assert e["emails"] == ["a@b.com"]
    assert e["ibans"] == ["DE89370400440532013000"]
    assert e["gstin"] == ["29ABCDE1234F1Z5"]
    assert e["card_numbers_masked"] == ["**** **** **** 1111"]
    r = redact(t)
    for secret in ("a@b.com", "4111", "DE89", "29ABCDE"):
        assert secret not in r


def test_key_values():
    kv = key_values("Invoice No: INV-1   Date: 2026-01-01\nCustomer ID      C-77")
    assert {"key": "Invoice No", "value": "INV-1"} in kv
    assert {"key": "Customer ID", "value": "C-77"} in kv


# ---------------------------------------------------------------- documents

def _check_invoice(r):
    f = r["fields"]
    assert r["doc_type"] == "invoice"
    assert f["invoice_number"] == "INV-2026-0042"
    assert f["invoice_date"] == "2026-03-15"
    assert f["due_date"] == "2026-04-14"
    assert f["vendor_name"] == "Northwind Supplies Ltd"
    assert f["currency"] == "GBP"
    assert f["subtotal"] == 288.0 and f["tax"] == 57.6 and f["total"] == 345.6
    assert f["totals_check"] == "ok"
    assert len(r["line_items"]) == 4
    assert r["line_items"][0] == {"description": "Espresso beans 1kg", "quantity": 4.0, "unit_price": 18.5, "amount": 74.0}


def test_digital_invoice(samples):
    r = extract_document(samples["invoice_pdf"].read_bytes(), "invoice.pdf")
    assert r["methods"] == ["text-layer"]
    _check_invoice(r)
    # bundled template auto-applies
    assert r["template"] == "northwind-invoice" or r["template"] is None


@needs_ocr
def test_scanned_invoice(samples):
    r = extract_document(samples["invoice_scan"].read_bytes(), "invoice_scan.png")
    assert r["methods"][0].startswith("ocr:")
    _check_invoice(r)


@needs_ocr
def test_receipt(samples):
    r = extract_document(samples["receipt"].read_bytes(), "receipt.png")
    f = r["fields"]
    assert r["doc_type"] == "receipt"
    assert f["merchant"] == "GREEN LEAF MARKET"
    assert f["date"] == "2026-09-12"  # day-first inferred from INR
    assert f["total"] == 934.5 and f["currency"] == "INR"
    assert f["payment_method"] == "UPI"
    assert len(r["line_items"]) == 4


def test_bank_statement(samples):
    r = extract_document(samples["statement"].read_bytes(), "statement.pdf")
    f = r["fields"]
    assert r["doc_type"] == "bank_statement"
    assert f["opening_balance"] == 2450.0 and f["closing_balance"] == 3967.72
    assert f["period_start"] == "2026-09-01" and f["period_end"] == "2026-09-30"
    tx = r["line_items"]
    assert len(tx) == 5
    assert tx[0]["credit"] == 3200.0 and tx[1]["debit"] == 1400.0 and tx[-1]["balance"] == 3967.72


def test_inline_template(samples):
    tpl = {"name": "t", "fields": {"po": {"regex": r"PO-(\d+)"}, "due": {"label": "Total Due", "type": "amount"},
                                   "has_iban": {"regex": "IBAN", "type": "bool"}}}
    r = extract_document(samples["invoice_pdf"].read_bytes(), "invoice.pdf", template=tpl)
    assert r["fields"]["po"] == "7781" and r["fields"]["due"] == 345.6 and r["fields"]["has_iban"] is True


def test_redact_flag(samples):
    r = extract_document(samples["invoice_pdf"].read_bytes(), "invoice.pdf", redact=True)
    assert "billing@northwind.example" not in r["text"]
    assert r["fields"]["email"] == "[REDACTED]"


def test_exports(samples):
    r = extract_document(samples["invoice_pdf"].read_bytes(), "invoice.pdf")
    assert json.loads(export(r, "json"))["fields"]["total"] == 345.6
    csv_text = export([r, r], "csv").decode("utf-8-sig")
    assert csv_text.splitlines()[0].startswith("file,doc_type")
    assert len(csv_text.strip().splitlines()) == 3
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(export([r], "xlsx")))
    assert wb.sheetnames[:2] == ["Documents", "Line items"]
    assert wb["Line items"].max_row == 5
    assert b"| invoice_number |" in export(r, "md")


# ---------------------------------------------------------------- API / CLI / MCP

def test_api(samples):
    from fastapi.testclient import TestClient

    from freeocr.api import app

    c = TestClient(app)
    assert c.get("/health").json()["status"] == "ok"
    assert "invoice" in c.get("/v1/info").json()["doc_types"]
    assert c.get("/").status_code == 200
    pdf = samples["invoice_pdf"].read_bytes()
    r = c.post("/v1/extract", files={"file": ("invoice.pdf", pdf, "application/pdf")})
    assert r.status_code == 200 and r.json()["fields"]["total"] == 345.6
    r = c.post("/v1/extract", files={"file": ("invoice.pdf", pdf)}, data={"format": "csv"})
    assert r.headers["content-type"].startswith("text/csv")
    r = c.post("/v1/batch", files=[("files", ("a.pdf", pdf)), ("files", ("b.pdf", samples["statement"].read_bytes()))])
    assert r.status_code == 200 and r.content[:2] == b"PK"  # xlsx zip
    r = c.post("/v1/batch", files=[("files", ("a.pdf", pdf))], data={"format": "json"})
    assert r.json()["count"] == 1
    r = c.post("/v1/export/csv", json=[{"file": "x", "doc_type": "invoice", "fields": {"total": 1}}])
    assert "total" in r.text
    r = c.post("/v1/ocr", files={"file": ("invoice.pdf", pdf)})
    assert "Northwind" in r.json()["text"]
    r = c.post("/v1/redact", files={"file": ("invoice.pdf", pdf)})
    assert "[EMAIL]" in r.json()["text"]
    r = c.post("/v1/extract", files={"file": ("empty.pdf", b"")})
    assert r.status_code == 422


def test_cli(samples, tmp_path):
    out = tmp_path / "out.xlsx"
    folder = samples["invoice_pdf"].parent
    p = subprocess.run([sys.executable, "-m", "freeocr", "extract", str(samples["invoice_pdf"]),
                        str(samples["statement"]), "-o", str(out)], capture_output=True, text=True, cwd=ROOT)
    assert p.returncode == 0, p.stderr
    assert out.exists() and out.stat().st_size > 1000
    w = tmp_path / "watch.csv"
    p = subprocess.run([sys.executable, "-m", "freeocr", "watch", str(folder), "-o", str(w), "--once",
                        "--engine", "auto"], capture_output=True, text=True, cwd=ROOT)
    assert p.returncode == 0, p.stderr
    assert w.read_text(encoding="utf-8-sig").count("\n") >= 3


def test_mcp(samples):
    msgs = [
        {"jsonrpc": "2.0", "id": 1, "method": "initialize", "params": {"protocolVersion": "2024-11-05"}},
        {"jsonrpc": "2.0", "method": "notifications/initialized"},
        {"jsonrpc": "2.0", "id": 2, "method": "tools/list"},
        {"jsonrpc": "2.0", "id": 3, "method": "tools/call",
         "params": {"name": "extract_document", "arguments": {"path": str(samples["invoice_pdf"])}}},
    ]
    p = subprocess.run([sys.executable, "-m", "freeocr", "mcp"], input="\n".join(map(json.dumps, msgs)) + "\n",
                       capture_output=True, text=True, cwd=ROOT, timeout=120)
    replies = [json.loads(x) for x in p.stdout.splitlines()]
    assert [r["id"] for r in replies] == [1, 2, 3]
    assert len(replies[1]["result"]["tools"]) == 4
    data = json.loads(replies[2]["result"]["content"][0]["text"])
    assert data["fields"]["total"] == 345.6


def test_llm_assist(samples, monkeypatch):
    import freeocr.llm as llm

    class Resp:
        def raise_for_status(self):
            pass

        def json(self):
            return {"choices": [{"message": {"content": 'Sure: {"store_code": "NW-01", "invoice_number": "x"}'}}]}

    seen = {}

    def fake_post(url, json, headers, timeout):
        seen["url"], seen["body"] = url, json
        return Resp()

    monkeypatch.setenv("FREEOCR_LLM_URL", "http://localhost:9/v1")
    monkeypatch.setenv("FREEOCR_LLM_MODEL", "local-model")
    monkeypatch.setattr(llm.httpx, "post", fake_post)
    r = extract_document(samples["invoice_pdf"].read_bytes(), "invoice.pdf",
                         fields=["store_code", "invoice_number"], use_llm=True)
    assert seen["url"] == "http://localhost:9/v1/chat/completions"
    assert r["fields"]["store_code"] == "NW-01"
    assert r["fields"]["invoice_number"] == "INV-2026-0042"  # rules win, model only fills gaps


# ---------------------------------------------------------------- searchable PDF

@pytest.mark.skipif(not __import__("freeocr.ocr", fromlist=["x"]).tesseract_available(), reason="needs tesseract")
def test_searchable_pdf(samples, tmp_path):
    import pdfplumber

    from freeocr.searchable import make_searchable_pdf

    pdf, stats = make_searchable_pdf(samples["invoice_scan"].read_bytes(), "scan.png")
    assert pdf[:5] == b"%PDF-" and stats == {"pages": 1, "ocr_pages": 1, "kept_pages": 0}
    with pdfplumber.open(io.BytesIO(pdf)) as d:
        assert "Northwind" in (d.pages[0].extract_text() or "")
    # digital pages are copied through, not re-OCR'd
    _, stats = make_searchable_pdf(samples["invoice_pdf"].read_bytes(), "invoice.pdf")
    assert stats["kept_pages"] == 1 and stats["ocr_pages"] == 0

    from fastapi.testclient import TestClient

    from freeocr.api import app

    r = TestClient(app).post("/v1/searchable-pdf", files={"file": ("scan.png", samples["receipt"].read_bytes())})
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf"
    assert r.headers["x-ocr-pages"] == "1"

    out = tmp_path / "out.pdf"
    p = subprocess.run([sys.executable, "-m", "freeocr", "pdf", str(samples["receipt"]), "-o", str(out)],
                       capture_output=True, text=True, cwd=ROOT)
    assert p.returncode == 0, p.stderr
    assert out.read_bytes()[:5] == b"%PDF-"
