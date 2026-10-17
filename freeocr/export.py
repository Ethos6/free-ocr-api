"""Export results as JSON, CSV, Excel (XLSX), Markdown or plain text.

Batch exports put one document per row (perfect for data-entry), and line items /
transactions on a second sheet so nothing gets lost.
"""

from __future__ import annotations

import csv
import io
import json
from typing import Any

FORMATS = ("json", "csv", "xlsx", "md", "txt")
MEDIA_TYPES = {
    "json": "application/json",
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "md": "text/markdown; charset=utf-8",
    "txt": "text/plain; charset=utf-8",
}


def _cell(v: Any) -> Any:
    if isinstance(v, (list, tuple)):
        return "; ".join(str(x) for x in v)
    if isinstance(v, dict):
        return json.dumps(v, ensure_ascii=False)
    return v


def summary_rows(results: list[dict]) -> tuple[list[str], list[dict]]:
    base = ["file", "doc_type", "pages", "ocr_confidence"]
    keys: list[str] = []
    rows = []
    for r in results:
        row = {"file": r.get("file"), "doc_type": r.get("doc_type"), "pages": r.get("page_count"),
               "ocr_confidence": r.get("ocr_confidence")}
        if r.get("error"):
            row["error"] = r["error"]
        for k, v in (r.get("fields") or {}).items():
            row[k] = _cell(v)
            if k not in keys:
                keys.append(k)
        rows.append(row)
    cols = base + keys + (["error"] if any("error" in r for r in rows) else [])
    return cols, rows


def item_rows(results: list[dict]) -> tuple[list[str], list[dict]]:
    keys: list[str] = []
    rows = []
    for r in results:
        for i, it in enumerate(r.get("line_items") or [], 1):
            row = {"file": r.get("file"), "row": i, **{k: _cell(v) for k, v in it.items()}}
            for k in it:
                if k not in keys:
                    keys.append(k)
            rows.append(row)
    return ["file", "row"] + keys, rows


def to_csv(results: list[dict], items: bool = False) -> bytes:
    cols, rows = item_rows(results) if items else summary_rows(results)
    if not items and len(results) == 1 and not results[0].get("fields") and results[0].get("line_items"):
        cols, rows = item_rows(results)
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=cols, extrasaction="ignore")
    w.writeheader()
    w.writerows(rows)
    return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM so Excel opens UTF-8 correctly


def to_xlsx(results: list[dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="2F5496")

    def sheet(ws, cols, rows):
        ws.append(cols)
        for cell in ws[1]:
            cell.font, cell.fill = head_font, head_fill
        for r in rows:
            ws.append([r.get(c) for c in cols])
        for i, col in enumerate(cols, 1):
            width = max([len(str(col))] + [len(str(r.get(col) or "")) for r in rows[:500]])
            ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, width + 2))
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = ws.dimensions

    ws = wb.active
    ws.title = "Documents"
    sheet(ws, *summary_rows(results))
    cols, rows = item_rows(results)
    if rows:
        sheet(wb.create_sheet("Line items"), cols, rows)
    tables = [(r.get("file"), t) for r in results for t in (r.get("tables") or [])]
    for n, (fname, t) in enumerate(tables[:20], 1):
        ws = wb.create_sheet(f"Table {n}")
        ws.append([f"Source: {fname}"])
        for row in t:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_markdown(results: list[dict]) -> bytes:
    out = []
    for r in results:
        out.append(f"## {r.get('file') or 'document'}  ({r.get('doc_type')})\n")
        if r.get("fields"):
            out.append("| Field | Value |\n|---|---|")
            out += [f"| {k} | {_cell(v)} |" for k, v in r["fields"].items()]
            out.append("")
        items = r.get("line_items") or []
        if items:
            keys = list(dict.fromkeys(k for it in items for k in it))
            out.append("| " + " | ".join(keys) + " |")
            out.append("|" + "---|" * len(keys))
            out += ["| " + " | ".join(str(_cell(it.get(k, ""))) for k in keys) + " |" for it in items]
            out.append("")
    return "\n".join(out).encode("utf-8")


def to_text(results: list[dict]) -> bytes:
    return "\n\n".join(f"===== {r.get('file')} =====\n{r.get('text', '')}" for r in results).encode("utf-8")


def export(results: list[dict] | dict, fmt: str = "json") -> bytes:
    if isinstance(results, dict):
        results = [results]
    fmt = fmt.lower()
    if fmt == "json":
        data = results[0] if len(results) == 1 else results
        return json.dumps(data, ensure_ascii=False, indent=2, default=str).encode("utf-8")
    if fmt == "csv":
        return to_csv(results)
    if fmt == "xlsx":
        return to_xlsx(results)
    if fmt == "md":
        return to_markdown(results)
    if fmt == "txt":
        return to_text(results)
    raise ValueError(f"unknown format '{fmt}', use one of {FORMATS}")
